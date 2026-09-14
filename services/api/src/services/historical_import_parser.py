"""Parser seguro e determinístico para fontes CSV/XLSX de importação histórica."""
from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import zipfile
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

from src.services.csv_import_service import HEADER_ALIASES, normalize_header

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_DATA_ROWS = 100_000
MAX_CELLS = 1_000_000
MAX_COLUMNS = 100
PREVIEW_ROWS = 100
SUPPORTED_FORMATS = {"csv", "xlsx"}
_CSV_MIME_TYPES = frozenset({"text/csv", "application/csv", "application/vnd.ms-excel"})
_XLSX_MIME_TYPES = frozenset({"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
SUPPORTED_FIELDS = {
    "name", "website", "phone", "whatsapp", "email", "city", "state",
    "address", "cnpj", "category", "contact_name", "linkedin", "instagram",
}


class ImportParseError(ValueError):
    """Erro de validação segura da origem, sem efeitos no CRM."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class ParsedSource:
    source_hash: str
    source_format: str
    headers: list[str]
    rows: list[list[str]]
    preview_rows: list[list[str]]
    suggested_mapping: dict[str, str | None]
    mapping_version: str


def _reject_formula(value: Any) -> None:
    if not isinstance(value, str):
        return
    stripped = value.lstrip()
    if stripped and stripped[0] in "=+-@":
        raise ImportParseError("FORMULA_NOT_ALLOWED", "A origem contém uma fórmula ou expressão não permitida.")


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    _reject_formula(value)
    if hasattr(value, "isoformat") and not isinstance(value, str):
        return value.isoformat()
    return str(value).strip()


def _validate_dimensions(rows: list[list[str]], headers: list[str]) -> None:
    if not headers or not any(item.strip() for item in headers):
        raise ImportParseError("HEADER_REQUIRED", "O arquivo precisa conter uma linha de cabeçalho.")
    if not rows or not any(any(item.strip() for item in row) for row in rows):
        raise ImportParseError("DATA_REQUIRED", "O arquivo precisa conter pelo menos uma linha de dados.")
    if len(headers) > MAX_COLUMNS:
        raise ImportParseError("COLUMN_LIMIT", f"O arquivo excede o limite de {MAX_COLUMNS} colunas.")
    if len(rows) > MAX_DATA_ROWS:
        raise ImportParseError("ROW_LIMIT", f"O arquivo excede o limite de {MAX_DATA_ROWS} linhas.")
    if len(headers) * (len(rows) + 1) > MAX_CELLS:
        raise ImportParseError("CELL_LIMIT", f"O arquivo excede o limite de {MAX_CELLS} células.")
    normalized = [item.strip().lower() for item in headers]
    if any(not item for item in normalized):
        raise ImportParseError("HEADER_INVALID", "Todas as colunas precisam ter nome.")
    if len(set(normalized)) != len(normalized):
        raise ImportParseError("HEADER_DUPLICATE", "O arquivo contém nomes de coluna duplicados.")


def _header_index(rows: list[list[str]]) -> int:
    for index, row in enumerate(rows[:15]):
        known = sum(1 for cell in row if normalize_header(cell) in HEADER_ALIASES)
        if known >= 2:
            return index
    return 0


def _suggest_mapping(headers: list[str]) -> dict[str, str | None]:
    result: dict[str, str | None] = {}
    for header in headers:
        candidate = normalize_header(header)
        result[header] = candidate if candidate in SUPPORTED_FIELDS else None
    return result


def mapping_version(headers: list[str], mapping: dict[str, str | None]) -> str:
    payload = json.dumps({"headers": headers, "mapping": mapping}, ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:32]


def validate_mapping(headers: list[str], mapping: dict[str, str | None]) -> dict[str, str | None]:
    if not isinstance(mapping, dict) or set(mapping) != set(headers):
        raise ImportParseError("MAPPING_INCOMPLETE", "Cada coluna da origem deve ser mapeada ou explicitamente ignorada.")
    normalized: dict[str, str | None] = {}
    used: set[str] = set()
    for header in headers:
        target = mapping.get(header)
        if target in (None, "", "ignore", "ignored"):
            normalized[header] = None
            continue
        if target not in SUPPORTED_FIELDS:
            raise ImportParseError("MAPPING_FIELD_INVALID", f"Campo de destino inválido: {target}.")
        if target in used:
            raise ImportParseError("MAPPING_FIELD_DUPLICATE", f"O campo {target} recebeu mais de uma coluna.")
        used.add(target)
        normalized[header] = target
    if "name" not in used:
        raise ImportParseError("NAME_REQUIRED", "O mapping precisa indicar a coluna de nome da empresa.")
    return normalized


def _safe_preview(rows: list[list[str]]) -> list[list[str]]:
    return [[html.escape(value, quote=True) for value in row] for row in rows[:PREVIEW_ROWS]]


def _declared_extension(filename: str | None) -> str:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        return "xlsx"
    if name.endswith(".csv"):
        return "csv"
    raise ImportParseError("FORMAT_INVALID", "Apenas arquivos CSV ou XLSX são aceitos.")


def _validate_declared_type(content: bytes, source_format: str, content_type: str | None) -> None:
    mime = (content_type or "").split(";", 1)[0].strip().lower()
    allowed = _XLSX_MIME_TYPES if source_format == "xlsx" else _CSV_MIME_TYPES
    if mime and mime not in allowed:
        raise ImportParseError("MIME_INVALID", "O tipo declarado não corresponde ao formato do arquivo.")
    is_zip = zipfile.is_zipfile(io.BytesIO(content))
    if source_format == "xlsx":
        if not is_zip or not content.startswith(b"PK"):
            raise ImportParseError("MIME_INVALID", "A assinatura do arquivo XLSX é inválida.")
    elif is_zip:
        raise ImportParseError("MIME_INVALID", "A extensão CSV não corresponde ao conteúdo do arquivo.")


def _csv_delimiter(text: str) -> str:
    sample = text[:4096]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        comma_count = sample.count(",")
        semicolon_count = sample.count(";")
        tab_count = sample.count("\t")
        if semicolon_count > comma_count and semicolon_count >= tab_count:
            return ";"
        if tab_count > comma_count and tab_count > semicolon_count:
            return "\t"
        return ","


def parse_source(content: bytes, filename: str | None, content_type: str | None = None) -> ParsedSource:
    if len(content) > MAX_FILE_BYTES:
        raise ImportParseError("FILE_TOO_LARGE", f"O arquivo excede o limite de {MAX_FILE_BYTES // (1024 * 1024)} MiB.")
    source_format = _declared_extension(filename)
    _validate_declared_type(content, source_format, content_type)
    if source_format == "xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if archive.testzip() is not None:
                    raise ImportParseError("XLSX_INVALID", "O arquivo XLSX está corrompido.")
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise ImportParseError("XLSX_TOO_LARGE", "O conteúdo descompactado excede o limite de segurança.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            sheet = workbook.active
            raw_rows = []
            total_cells = 0
            for row in sheet.iter_rows(values_only=True):
                if len(row) > MAX_COLUMNS:
                    raise ImportParseError("COLUMN_LIMIT", f"O arquivo excede o limite de {MAX_COLUMNS} colunas.")
                total_cells += len(row)
                if total_cells > MAX_CELLS:
                    raise ImportParseError("CELL_LIMIT", f"O arquivo excede o limite de {MAX_CELLS} células.")
                if len(raw_rows) > MAX_DATA_ROWS + 15:
                    raise ImportParseError("ROW_LIMIT", f"O arquivo excede o limite de {MAX_DATA_ROWS} linhas.")
                raw_rows.append([_cell_to_text(value) for value in row])
            workbook.close()
        except ImportParseError:
            raise
        except Exception as exc:
            raise ImportParseError("XLSX_INVALID", "Não foi possível ler o arquivo XLSX.") from exc
        source_format = "xlsx"
    else:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ImportParseError("ENCODING_INVALID", "CSV deve estar em UTF-8 ou UTF-8 com BOM.") from exc
        delimiter = _csv_delimiter(text)
        try:
            raw_rows = []
            for row in csv.reader(io.StringIO(text), delimiter=delimiter, strict=True):
                if len(row) > MAX_COLUMNS:
                    raise ImportParseError("COLUMN_LIMIT", f"O arquivo excede o limite de {MAX_COLUMNS} colunas.")
                raw_rows.append([_cell_to_text(value) for value in row])
        except csv.Error as exc:
            raise ImportParseError("CSV_INVALID", "O arquivo CSV está malformado.") from exc

    if not raw_rows:
        raise ImportParseError("EMPTY_FILE", "O arquivo não contém linhas.")
    header_idx = _header_index(raw_rows)
    headers = raw_rows[header_idx]
    rows = []
    for row in raw_rows[header_idx + 1:]:
        padded = row[:MAX_COLUMNS] + [""] * max(0, len(headers) - len(row))
        rows.append(padded[:len(headers)])
    _validate_dimensions(rows, headers)
    suggested = _suggest_mapping(headers)
    version = mapping_version(headers, suggested)
    return ParsedSource(
        source_hash=hashlib.sha256(content).hexdigest(),
        source_format=source_format,
        headers=headers,
        rows=rows,
        preview_rows=_safe_preview(rows),
        suggested_mapping=suggested,
        mapping_version=version,
    )
